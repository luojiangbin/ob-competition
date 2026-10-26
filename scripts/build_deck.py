"""生成答辩 PPT：docs/答辩.pptx（10 页，16:9）"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn

NAVY = RGBColor(0x0B, 0x1F, 0x33)
ACCENT = RGBColor(0x1F, 0x8A, 0x8C)
GREY = RGBColor(0x5A, 0x6B, 0x7B)
LIGHT = RGBColor(0xF3, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x39, 0x2B)

FONT = "微软雅黑"


def set_font(run, size: int, bold: bool = False, color: RGBColor = NAVY) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", FONT)


def textbox(slide, left, top, width, height, text, size=18, bold=False,
            color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = Inches(0.05)
    frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    set_font(run, size, bold, color)
    return box


def bullets(slide, left, top, width, height, items, size=17, gap=10, color=NAVY):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(gap)
        run = paragraph.add_run()
        run.text = item
        set_font(run, size, False, color)
    return box


def flat(shape):
    """去掉主题默认的阴影与描边样式。"""
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def band(slide, top, height, color=LIGHT):
    shape = slide.shapes.add_shape(1, Inches(0), Inches(top), Inches(13.333), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return flat(shape)


def header(slide, title, kicker=""):
    textbox(slide, 0.7, 0.45, 11.9, 0.7, title, 30, True)
    if kicker:
        textbox(slide, 0.72, 1.15, 11.9, 0.4, kicker, 14, False, GREY)
    line = slide.shapes.add_shape(1, Inches(0.7), Inches(1.62), Inches(0.9), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def table(slide, left, top, width, rows, col_widths, font_size=14, header_fill=NAVY):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top),
                                   Inches(width), Inches(0.45 * len(rows)))
    tbl = shape.table
    for index, width_in in enumerate(col_widths):
        tbl.columns[index].width = Inches(width_in)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = cell.text_frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = str(value)
            if r == 0:
                set_font(run, font_size, True, WHITE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
            else:
                set_font(run, font_size, False, NAVY)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
    return tbl


def new_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY
        shape.line.fill.background()
    return slide


def build(out: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 封面
    slide = new_slide(prs, dark=True)
    textbox(slide, 1.0, 1.9, 11.3, 0.5, "2026 第六届 OceanBase 数据库大赛 · PowerContext 方向", 16, False, RGBColor(0x9F, 0xC5, 0xC6))
    textbox(slide, 1.0, 2.5, 11.3, 1.2, "课程项目上下文接力", 52, True, WHITE)
    textbox(slide, 1.0, 3.75, 11.0, 0.6, "让课设上下文活过下一次会话、下一次交接、下一台电脑", 22, False, RGBColor(0xD8, 0xE2, 0xE8))
    textbox(slide, 1.0, 5.4, 11.0, 1.2, "队伍 10086 ｜ 罗江斌 · 宜宾学院 ｜ 指导老师：薛一兰", 18, False, WHITE)

    # 2 问题
    slide = new_slide(prs)
    header(slide, "一个学生的真实一周", "5–6 门课，每门一个课设，上下文活不过一次会话")
    data = [
        ("会话断裂", "上周敲定的架构，新开会话全忘了，要从头解释"),
        ("交接断裂", "队友接手，得把要求、选型、踩过的坑重讲一遍"),
        ("工具断裂", "换个 Agent、换台电脑，上下文带不过去"),
    ]
    top = 2.0
    for name, desc in data:
        band(slide, top, 1.15)
        textbox(slide, 0.85, top + 0.18, 3.0, 0.6, name, 20, True, ACCENT)
        textbox(slide, 4.0, top + 0.2, 8.4, 0.7, desc, 17, False, NAVY)
        top += 1.35
    textbox(slide, 0.85, 6.3, 11.8, 0.5, "真正需要延续的是决策、约束、进度——不是聊天文本。", 18, True, NAVY)

    # 3 现有做法
    slide = new_slide(prs)
    header(slide, "现有做法为什么都不够", "")
    table(
        slide, 0.75, 1.95, 11.9,
        [
            ["做法", "问题"],
            ["重读聊天记录", "上下文窗口有限，跨会话根本不存在"],
            ["RAG 检索聊天记录", "检索到的是原话碎片，不是结论"],
            ["本地 memory 文件", "绑在单个 Agent 上，换工具就失效"],
            ["让人口头复述", "成本最高，且必然遗漏"],
        ],
        [4.0, 7.9],
        font_size=16,
    )
    textbox(slide, 0.75, 5.25, 11.9, 0.5, "结论：不存对话，存结论。", 20, True, NAVY)

    # 4 方案
    slide = new_slide(prs)
    header(slide, "方案：一门课 = 一个隔离空间", "四类上下文各归其位")
    cards = [
        ("Scope", "一门课", "整学期隔离，不串数据"),
        ("Memory", "约束 / 决策 / 踩坑 / 下一步", "跨会话长期有效"),
        ("Handoff", "队友交接", "契约 → 临时包 → 固化 → 回执"),
        ("Experience", "踩坑经验", "审核通过才可复用"),
    ]
    left = 0.75
    for name, mid, desc in cards:
        card = slide.shapes.add_shape(1, Inches(left), Inches(1.95), Inches(2.85), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.fill.background()
        flat(card)
        head = slide.shapes.add_shape(1, Inches(left), Inches(1.95), Inches(2.85), Inches(0.55))
        head.fill.solid()
        head.fill.fore_color.rgb = ACCENT
        head.line.fill.background()
        flat(head)
        textbox(slide, left + 0.18, 2.03, 2.5, 0.4, name, 18, True, WHITE)
        textbox(slide, left + 0.18, 2.62, 2.5, 0.5, mid, 14, True, NAVY)
        textbox(slide, left + 0.18, 3.22, 2.5, 0.9, desc, 13, False, GREY)
        left += 3.02
    textbox(slide, 0.75, 5.15, 11.9, 0.5,
            "四类上下文 ↔ 四个原语，一一对应，不多不少。", 18, True, NAVY)

    # 5 演示
    slide = new_slide(prs)
    header(slide, "现场演示：新会话里的第一次提问", "")
    textbox(slide, 0.9, 2.0, 11.6, 0.5, "问：课设允许用 ORM 吗？", 22, True, NAVY)
    textbox(slide, 0.9, 2.7, 11.6, 0.5, "召回：课设不允许使用 ORM，必须手写 SQL。", 22, True, ACCENT)
    textbox(slide, 0.9, 3.45, 11.6, 0.9,
            "问的是「允许吗」，存的是「不允许」——字面检索原本匹配不上。\n把记忆改写成提问者会用的措辞，命中率从 80% 提到 100%。",
            15, False, GREY)
    holder = slide.shapes.add_shape(1, Inches(0.9), Inches(4.55), Inches(11.55), Inches(2.3))
    holder.fill.solid()
    holder.fill.fore_color.rgb = LIGHT
    holder.line.color.rgb = RGBColor(0xC9, 0xD4, 0xDC)
    holder.line.width = Pt(1)
    flat(holder)
    textbox(slide, 0.9, 5.5, 11.55, 0.5, "（此处放现场录屏截图）", 14, False, GREY, PP_ALIGN.CENTER)

    # 6 数据
    slide = new_slide(prs)
    header(slide, "实验：五组对照，全部可复现", "判定方式为「上下文是否命中期望关键词」，不依赖任何大模型")
    table(
        slide, 0.75, 1.95, 11.9,
        [
            ["组", "条件", "准确率"],
            ["A", "自然写法", "80%"],
            ["B", "写入改写（面向提问）", "100%"],
            ["C", "留出集 10 题（未参与调优）", "100%"],
            ["D", "同义词集（刻意换词）", "10%"],
            ["E", "同义词集 + 查询改写", "90%"],
        ],
        [1.2, 8.4, 2.3],
        font_size=16,
    )
    textbox(slide, 0.75, 5.15, 11.9, 0.5, "D 组是核心发现：纯字面检索的天花板在语义层。", 18, True, RED)
    textbox(slide, 0.75, 5.75, 11.9, 0.5, "复现：cd experiments && python evaluate.py --cases cases_synonym.jsonl --expand", 14, False, GREY)

    # 7 失败分析
    slide = new_slide(prs)
    header(slide, "失败模式：一个字之差就整条漏掉", "")
    table(
        slide, 0.75, 1.9, 11.9,
        [
            ["提问", "记忆里的写法", "结果"],
            ["课设允许用 ORM 吗？", "禁止使用 ORM", "未召回"],
            ["为什么用深搜？", "DFS", "未召回"],
            ["进程控制块怎么存的？", "PCB", "未召回"],
        ],
        [4.6, 5.0, 2.3],
        font_size=15,
    )
    bullets(slide, 0.75, 4.15, 11.9, 2.4, [
        "▸ 全文检索是字面匹配：允许 / 禁止、深搜 / DFS，对不上就漏",
        "▸ 写入时改写记忆 → 零成本，+20pt",
        "▸ 查询时同义词扩展 → 课程域可行，10% → 90%",
        "▸ 真正的解是向量 / 混合检索，这是下一步",
    ], size=17, gap=12)

    # 8 边界
    slide = new_slide(prs)
    header(slide, "我们明确不声称", "主动划边界，比被问穿强")
    bullets(slide, 0.72, 2.1, 11.8, 3.6, [
        "不声称能处理开放域问答——只针对课程项目这个封闭域",
        "不声称纯 FTS 够用——D 组 10% 就是反证",
        "不声称 100% 是真实水平——C 组留出题仍与记忆共享关键词",
        "未审核的经验不进入召回；模型提议不能自动成为团队共识",
    ], size=19, gap=18)
    textbox(slide, 0.72, 5.6, 11.8, 0.5, "宁可把边界说清楚，也不把数字做大。", 20, True, ACCENT)

    # 9 能力使用
    slide = new_slide(prs)
    header(slide, "PowerContext 能力使用：7 个用了 5 个", "不吹牛，用到的才写")
    table(
        slide, 0.75, 1.95, 11.9,
        [
            ["能力", "用途", "状态"],
            ["Scope", "一门课一个隔离空间", "已实现"],
            ["Atomic Memory", "约束 / 决策 / 踩坑 / 下一步", "已实现"],
            ["PreparedContext", "召回主链路", "已实现"],
            ["Handoff", "队友交接四步流程", "已实现"],
            ["Experience", "踩坑 → 审核 → 复用", "已实现"],
            ["Skill / Profile", "套路导出、学生画像", "规划中"],
        ],
        [3.6, 6.6, 1.7],
        font_size=14,
    )
    textbox(slide, 0.75, 5.55, 11.9, 0.5, "没用到的两个写在规划里，不吹。", 18, True, ACCENT)

    # 10 结尾
    slide = new_slide(prs, dark=True)
    textbox(slide, 1.0, 2.4, 11.3, 1.2, "换设备、换工具、换队友", 40, True, WHITE)
    textbox(slide, 1.0, 3.5, 11.3, 1.2, "都不用再解释一遍", 40, True, RGBColor(0x9F, 0xC5, 0xC6))
    textbox(slide, 1.0, 5.2, 11.3, 0.6, "项目开源 · 一条命令可复现 · 数据可当场重跑", 18, False, RGBColor(0xD8, 0xE2, 0xE8))

    prs.save(out)
    print(f"saved {out} ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    target = Path(__file__).resolve().parent.parent / "docs" / "答辩.pptx"
    build(target)
